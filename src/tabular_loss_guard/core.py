"""Strict positional contracts: uncertainty must never become a pass."""
import csv
import hashlib
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class Unverified(ValueError):
    """The input lies outside the supported alignment/value contract."""


def source_rows(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        if sheet is not None:
            raise Unverified("--sheet applies only to XLSX")
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = [[(v, "text") for v in row] for row in csv.reader(stream)]
        return path.stem, rows
    if path.suffix.lower() != ".xlsx":
        raise Unverified("Only .csv and .xlsx sources are supported")
    book = load_workbook(path, data_only=False, read_only=False)
    try:
        if sheet is None and len(book.sheetnames) != 1:
            raise Unverified("Multiple sheets: select one with --sheet")
        if sheet is not None and sheet not in book.sheetnames:
            raise Unverified("Selected sheet does not exist")
        ws = book[sheet] if sheet else book.active
        if ws.merged_cells.ranges:
            raise Unverified("Merged cells have no supported positional contract")
        rows = []
        for row in ws:
            values = []
            for cell in row:
                value = cell.value
                kind = "text"
                if cell.data_type == "f" or isinstance(value, (date, datetime, time)):
                    kind = "unsupported"
                elif value is None:
                    value, kind = "", "blank"
                elif isinstance(value, bool):
                    value = "True" if value else "False"
                elif cell.data_type == "n":
                    kind = "number"
                values.append((str(value), kind))
            rows.append(values)
        return ws.title, rows
    finally:
        book.close()


def split_row(line):
    if not line.startswith("|") or not line.endswith("|"):
        raise Unverified("Tables must have outer pipes")
    # Only escaped pipes and escaped backslashes are decoded. Other Markdown
    # syntax remains unsupported rather than silently normalized.
    cells, current, i = [], "", 1
    while i < len(line) - 1:
        char = line[i]
        if char == "\\" and i + 1 < len(line) - 1 and line[i + 1] in "|\\":
            current += line[i + 1]
            i += 2
            continue
        if char == "|":
            cells.append(current.strip())
            current = ""
        else:
            current += char
        i += 1
    cells.append(current.strip())
    return cells


def markdown_rows(text):
    # Require exactly one contiguous pipe table. Headings/prose outside it are OK.
    blocks, block = [], []
    fenced = False
    for raw in text.splitlines():
        line = raw.strip()
        if line.startswith(("```", "~~~")):
            fenced = not fenced
            if block:
                blocks.append(block)
                block = []
            continue
        if fenced:
            continue
        if line.startswith("|"):
            block.append(split_row(line))
        elif block:
            blocks.append(block)
            block = []
    if block:
        blocks.append(block)
    if fenced or len(blocks) != 1:
        raise Unverified("Expected exactly one unfenced pipe table")
    rows = blocks[0]
    if len(rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
        raise Unverified("Missing Markdown header separator")
    if any(len(row) != len(rows[0]) for row in rows):
        raise Unverified("Ragged Markdown table")
    return [rows[0], *rows[2:]]


def audit(source, markdown, sheet=None):
    source, markdown = Path(source), Path(markdown)
    report = {"schema_version": 1, "policy": "ordered-stored-values-v1",
              "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
              "markdown_sha256": hashlib.sha256(markdown.read_bytes()).hexdigest(),
              "findings": []}
    try:
        name, expected = source_rows(source, sheet)
        observed = markdown_rows(markdown.read_text(encoding="utf-8-sig"))
        report["sheet"] = name
        if not expected or not expected[0] or any(len(r) != len(expected[0]) for r in expected):
            raise Unverified("Empty or ragged source table")
        if any(k == "unsupported" or v != v.strip() or "\n" in v or "\r" in v
               for row in expected for v, k in row):
            raise Unverified("Formulas, dates, multiline or edge-whitespace cells need an explicit policy")
        if any(re.search(r"[`*_<>&\[\]]", v) for row in observed for v in row):
            raise Unverified("Inline Markdown/HTML is outside the plain-cell contract")
        if len(expected) != len(observed) or len(expected[0]) != len(observed[0]):
            report["findings"].append({"reason": "shape_changed", "expected": [len(expected), len(expected[0])],
                                       "observed": [len(observed), len(observed[0])]})
        for r, row in enumerate(expected):
            for c, (value, kind) in enumerate(row):
                actual = observed[r][c] if r < len(observed) and c < len(observed[r]) else None
                equal = value == actual
                if kind == "number" and actual is not None:
                    try:
                        numeric = Decimal(actual)
                        equal = bool(re.fullmatch(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?", actual)) and numeric.is_finite() and Decimal(value) == numeric
                    except InvalidOperation:
                        equal = False
                if not equal:
                    report["findings"].append({"cell": f"{get_column_letter(c + 1)}{r + 1}",
                                               "row": r + 1, "column": c + 1, "kind": kind,
                                               "expected": value, "observed": actual,
                                               "reason": "value_changed"})
        report["status"] = "FAIL" if report["findings"] else "PASS"
    except Unverified as exc:
        report.update(status="UNVERIFIED", reason=str(exc))
    return report
