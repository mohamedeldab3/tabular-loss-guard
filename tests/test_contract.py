import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from tabular_loss_guard.core import audit, markdown_rows, Unverified


def check(tmp_path, rows, table):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    source = tmp_path / "source.xlsx"
    book.save(source)
    book.close()
    target = tmp_path / "converted.md"
    target.write_text(table, encoding="utf-8")
    return audit(source, target)


@pytest.mark.parametrize("value,rendered", [(42, "42.0"), (0.1, "1e-1"), (1e-10, "1.000000e-10"),
                                               (123456789.123, "123456789.123")])
def test_numeric_equivalence(tmp_path, value, rendered):
    assert check(tmp_path, [["Value"], [value]], f"| Value |\n| --- |\n| {rendered} |\n")["status"] == "PASS"


def test_precision_loss(tmp_path):
    result = check(tmp_path, [["Value"], [123456789.123]], "| Value |\n| --- |\n| 1.234568e+08 |")
    assert result["status"] == "FAIL"
    assert result["findings"][0]["cell"] == "A2"


@pytest.mark.parametrize("error", ["#DIV/0!", "#REF!", "#VALUE!", "#N/A", "#NUM!", "#NAME?", "#NULL!"])
def test_errors_vs_blanks(tmp_path, error):
    rows = [["Case", "Value"], ["error", error], ["blank", None]]
    good = f"| Case | Value |\n| --- | --- |\n| error | {error} |\n| blank | |"
    assert check(tmp_path, rows, good)["status"] == "PASS"
    bad = good.replace(error, "NaN")
    result = check(tmp_path, rows, bad)
    assert result["status"] == "FAIL"
    assert result["findings"][0]["expected"] == error


@pytest.mark.parametrize("value,rendered", [("00123", "123"), ("مرحبا", "ابحرم"), (True, "False")])
def test_typed_mutations(tmp_path, value, rendered):
    assert check(tmp_path, [["Value"], [value]], f"| Value |\n| --- |\n| {rendered} |")["status"] == "FAIL"


def test_escaped_pipe_unicode(tmp_path):
    result = check(tmp_path, [["Code", "Name"], ["00123", "عربي | English"]],
                   "| Code | Name |\n| --- | --- |\n| 00123 | عربي \\| English |")
    assert result["status"] == "PASS"


@pytest.mark.parametrize("body", ["| b | 2 |\n| a | 1 |", "| a | 2 |\n| b | 1 |", "| a | 1 |"])
def test_row_identity_and_shape(tmp_path, body):
    result = check(tmp_path, [["ID", "Value"], ["a", 1], ["b", 2]], "| ID | Value |\n| --- | --- |\n" + body)
    assert result["status"] == "FAIL"


@pytest.mark.parametrize("value", ["=1+1", "line\nbreak", " padded "])
def test_unsupported_source(tmp_path, value):
    assert check(tmp_path, [["Value"], [value]], "| Value |\n| --- |\n| 2 |")["status"] == "UNVERIFIED"


@pytest.mark.parametrize("text", ["no table", "| a |\n| b |", "| a |\n| --- |\n| b | c |",
                                  "| a |\n| --- |\n\n| b |\n| --- |"])
def test_ambiguous_markdown(text):
    with pytest.raises(Unverified):
        markdown_rows(text)


def test_inline_syntax_unverified(tmp_path):
    assert check(tmp_path, [["Value"], ["hello"]], "| Value |\n| --- |\n| **hello** |")["status"] == "UNVERIFIED"


def test_csv_text_identity(tmp_path):
    source, target = tmp_path / "source.csv", tmp_path / "target.md"
    source.write_text("ID,Value\n001,42\n", encoding="utf-8")
    target.write_text("| ID | Value |\n| --- | --- |\n| 001 | 42 |", encoding="utf-8")
    assert audit(source, target)["status"] == "PASS"
    target.write_text("| ID | Value |\n| --- | --- |\n| 1 | 42 |", encoding="utf-8")
    assert audit(source, target)["status"] == "FAIL"


def test_sheet_selection_and_merged(tmp_path):
    source, target = tmp_path / "source.xlsx", tmp_path / "target.md"
    book = Workbook()
    book.create_sheet("Other")
    book.active.append(["Value"])
    book.active.append([42])
    book.save(source)
    target.write_text("| Value |\n| --- |\n| 42 |", encoding="utf-8")
    assert audit(source, target)["status"] == "UNVERIFIED"
    assert audit(source, target, "Sheet")["status"] == "PASS"
    book.active.merge_cells("A2:B2")
    book.save(source)
    book.close()
    assert audit(source, target, "Sheet")["status"] == "UNVERIFIED"


def test_cli_reports_and_exit(tmp_path):
    source, target = tmp_path / "source.csv", tmp_path / "target.md"
    source.write_text("Value\n<script>alert(1)</script>\n", encoding="utf-8")
    target.write_text("| Value |\n| --- |\n| changed |", encoding="utf-8")
    report, page = tmp_path / "report.json", tmp_path / "report.html"
    result = subprocess.run([sys.executable, "-m", "tabular_loss_guard.cli", str(source), str(target),
                             "--json", str(report), "--html", str(page)], capture_output=True)
    assert result.returncode == 1, result.stderr
    assert json.loads(report.read_text())["status"] == "FAIL"
    assert "<script>" not in page.read_text()
    assert "&lt;script&gt;" in page.read_text()
    result = subprocess.run([sys.executable, "-m", "tabular_loss_guard.cli", str(source), str(target),
                             "--json", str(source)], capture_output=True)
    assert result.returncode == 2
    assert source.read_text().startswith("Value")


@pytest.mark.parametrize("rendered,status", [("4_2", "UNVERIFIED"), ("NaN", "FAIL"), ("Infinity", "FAIL"), ("42kg", "FAIL")])
def test_non_numeric_spellings_fail(tmp_path, rendered, status):
    assert check(tmp_path, [["Value"], [42]], f"| Value |\n| --- |\n| {rendered} |")["status"] == status
